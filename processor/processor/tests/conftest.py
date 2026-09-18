"""Fixture builders that generate small, real (not mocked) PDF/DOCX/PPTX
files on disk for tests to extract from. Building them programmatically
with the exact same libraries used to author real documents (reportlab for
PDF, python-docx/python-pptx for their own formats) keeps fixtures tiny,
deterministic, and honest about what a real file of that format looks
like — no hand-crafted binary blobs to maintain.
"""

from __future__ import annotations

import pytest


@pytest.fixture
def sample_pdf_path(tmp_path) -> str:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter

    path = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)

    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 750, "Chapter One")
    c.setFont("Helvetica", 11)
    c.drawString(72, 700, "This is the first paragraph of chapter one.")
    c.drawString(72, 680, "This is the second sentence continuing the paragraph.")
    c.showPage()

    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 750, "Chapter Two")
    c.setFont("Helvetica", 11)
    c.drawString(72, 700, "Second chapter content here.")
    c.showPage()

    c.save()
    return str(path)


@pytest.fixture
def blank_pdf_path(tmp_path) -> str:
    """A PDF with pages but no text at all — stands in for a scanned document."""
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter

    path = tmp_path / "blank.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.rect(100, 100, 200, 200)  # a shape, but no text
    c.showPage()
    c.save()
    return str(path)


@pytest.fixture
def sample_docx_path(tmp_path) -> str:
    from docx import Document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This is the first paragraph.")
    document.add_paragraph("Bullet one", style="List Bullet")
    document.add_paragraph("Bullet two", style="List Bullet")
    document.add_heading("Background", level=2)
    document.add_paragraph("More detail here.")
    document.save(str(path))
    return str(path)


@pytest.fixture
def sample_pptx_path(tmp_path) -> str:
    from pptx import Presentation

    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # "Title and Content"

    slide1 = presentation.slides.add_slide(layout)
    slide1.shapes.title.text = "Welcome"
    body1 = slide1.placeholders[1].text_frame
    body1.text = "First point"
    second_paragraph = body1.add_paragraph()
    second_paragraph.text = "Second point"
    second_paragraph.level = 1

    slide2 = presentation.slides.add_slide(layout)
    slide2.shapes.title.text = "Details"
    slide2.placeholders[1].text_frame.text = "Some detail text"

    presentation.save(str(path))
    return str(path)
