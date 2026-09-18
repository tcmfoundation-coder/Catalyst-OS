"""DOCX extraction via python-docx.

Unlike PDF, Word documents carry real structural metadata: every
paragraph has a style name ("Heading 1", "List Bullet", "Normal", ...).
We read that directly instead of guessing — which is why this extractor
is far more reliable than the PDF one. DOCX has no fixed "pages" at the
XML level (pagination only happens when something actually lays the
document out for print/screen), so block.page stays None throughout.
"""

from __future__ import annotations

import re

from .base import ExtractionError
from ..models import Block, BlockType, DocumentMetadata, NormalizedDocument, Section, SourceFormat

_HEADING_STYLE_RE = re.compile(r"^heading\s*(\d+)$", re.IGNORECASE)


def _heading_level(style_name: str) -> int | None:
    name = style_name.strip()
    if name.lower() == "title":
        return 1
    match = _HEADING_STYLE_RE.match(name)
    return int(match.group(1)) if match else None


def _list_level(paragraph) -> int:
    # python-docx doesn't expose numbering depth as a plain attribute; the
    # underlying XML does, via the paragraph's numbering properties. We
    # fall back to 1 rather than fail extraction if that XML is absent or
    # shaped differently than expected (e.g. a list style without explicit
    # numPr, which does happen).
    try:
        num_pr = paragraph._p.pPr.numPr
        return int(num_pr.ilvl.val) + 1
    except AttributeError:
        return 1


class DocxExtractor:
    def extract(self, file_path: str) -> NormalizedDocument:
        from docx import Document as DocxDocument

        try:
            document = DocxDocument(file_path)
        except Exception as exc:
            raise ExtractionError(f"Could not read DOCX: {exc}") from exc

        core_props = document.core_properties
        metadata = DocumentMetadata(
            source_format=SourceFormat.docx,
            title=(core_props.title or None),
            author=(core_props.author or None),
        )

        sections: list[Section] = []
        current_heading: str | None = None
        current_level: int | None = None
        current_blocks: list[Block] = []

        def flush() -> None:
            nonlocal current_heading, current_level, current_blocks
            if current_heading is not None or current_blocks:
                sections.append(
                    Section(heading=current_heading, heading_level=current_level, blocks=current_blocks)
                )
            current_heading = None
            current_level = None
            current_blocks = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            level = _heading_level(style_name)

            if level is not None:
                flush()
                current_heading = text
                current_level = level
                continue

            if "list" in style_name.lower():
                current_blocks.append(
                    Block(type=BlockType.list_item, text=text, level=_list_level(paragraph))
                )
            else:
                current_blocks.append(Block(type=BlockType.paragraph, text=text))

        flush()
        return NormalizedDocument(metadata=metadata, sections=sections)
