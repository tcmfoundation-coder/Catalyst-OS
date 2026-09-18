"""PPTX extraction via python-pptx.

Each slide is its own Section — that boundary is unambiguous, unlike PDF
pages or DOCX headings. The slide's title placeholder (if any) becomes the
section heading; every other text frame's paragraphs become blocks, using
python-pptx's real `paragraph.level` (outline/indent depth) to tell body
text from nested bullets.
"""

from __future__ import annotations

from .base import ExtractionError
from ..models import Block, BlockType, DocumentMetadata, NormalizedDocument, Section, SourceFormat


class PptxExtractor:
    def extract(self, file_path: str) -> NormalizedDocument:
        from pptx import Presentation

        try:
            presentation = Presentation(file_path)
        except Exception as exc:
            raise ExtractionError(f"Could not read PPTX: {exc}") from exc

        core_props = presentation.core_properties
        metadata = DocumentMetadata(
            source_format=SourceFormat.pptx,
            title=(core_props.title or None),
            author=(core_props.author or None),
            slide_count=len(presentation.slides),
        )

        sections: list[Section] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            title_id = title_shape.shape_id if title_shape is not None else None
            heading_text = None
            if title_shape is not None and title_shape.has_text_frame:
                heading_text = title_shape.text.strip() or None

            blocks: list[Block] = []
            for shape in slide.shapes:
                if title_id is not None and shape.shape_id == title_id:
                    continue
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if not text:
                        continue
                    is_list = paragraph.level > 0
                    blocks.append(
                        Block(
                            type=BlockType.list_item if is_list else BlockType.paragraph,
                            text=text,
                            level=paragraph.level if is_list else None,
                            slide=slide_number,
                        )
                    )

            sections.append(
                Section(
                    heading=heading_text,
                    heading_level=1 if heading_text else None,
                    blocks=blocks,
                )
            )

        return NormalizedDocument(metadata=metadata, sections=sections)
